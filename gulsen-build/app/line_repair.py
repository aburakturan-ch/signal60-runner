from __future__ import annotations

import os
import re
import shutil
import uuid
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List, Sequence, Set, Tuple

from docx import Document
from docx.text.paragraph import Paragraph


TERMINAL_RE = re.compile(r'[.!?…][\"”’»)]*$')
PAGE_NUMBER_RE = re.compile(r'^\s*\d{1,4}\s*$')
BULLET_RE = re.compile(r'^\s*(?:[-–—•▪◦*]|\d+[.)]|[A-Za-zÇĞİÖŞÜçğıöşü][.)])\s+')
HEADING_NUMBER_RE = re.compile(r'^\s*(?:[IVXLCDM]+|\d+(?:\.\d+)*)\s*[-–—.)]?\s+\S+', re.I)
INTRO_VERSE_RE = re.compile(r'\b(?:buyurur|buyurdu|buyuruyor)\s*[:,]?\s*$', re.I)


@dataclass
class RepairResult:
    source: Path
    output: Path
    repaired_breaks: int
    skipped: bool = False
    message: str = ""


def _normalized(text: str) -> str:
    return re.sub(r'\s+', ' ', text.strip())


def _is_page_number(text: str) -> bool:
    return bool(PAGE_NUMBER_RE.match(text or ''))


def _looks_like_heading(text: str) -> bool:
    t = _normalized(text)
    if not t or len(t) > 120:
        return False
    if BULLET_RE.match(t):
        return True
    if HEADING_NUMBER_RE.match(t) and len(t) < 80:
        return True
    letters = [c for c in t if c.isalpha()]
    if len(letters) >= 4:
        upper_ratio = sum(c.isupper() for c in letters) / len(letters)
        if upper_ratio > 0.78 and len(t) < 90:
            return True
    return False


def _ends_terminal(text: str) -> bool:
    return bool(TERMINAL_RE.search(text.rstrip()))


def _opening_balance(text: str) -> int:
    # Positive means there is an unmatched opener, which strongly suggests continuation.
    pairs = [('(', ')'), ('[', ']'), ('{', '}')]
    return sum(text.count(a) - text.count(b) for a, b in pairs)


def _starts_lowercase(text: str) -> bool:
    t = text.lstrip(' \t\"“”\'‘’([{«')
    if not t:
        return False
    return t[0].islower()


def _starts_uppercase(text: str) -> bool:
    t = text.lstrip(' \t\"“”\'‘’([{«')
    if not t:
        return False
    return t[0].isupper()


def _short_line(text: str) -> bool:
    return len(_normalized(text)) <= 58


def _paragraph_frequencies(paragraphs: Sequence[Paragraph]) -> Counter:
    return Counter(_normalized(p.text) for p in paragraphs if _normalized(p.text))


def _protected_poetry_nodes(paragraphs: Sequence[Paragraph]) -> Set[int]:
    """Protect verse/quotation blocks introduced by 'buyurur:' style lines.

    The source material frequently contains actual verse where each line is intentionally
    a separate paragraph. Those lines must not be joined merely because they lack a period.
    """
    protected: Set[int] = set()
    n = len(paragraphs)
    for i, p in enumerate(paragraphs[:-1]):
        intro = _normalized(p.text)
        if not INTRO_VERSE_RE.search(intro):
            continue
        j = i + 1
        # Skip empty paragraphs immediately after the intro.
        while j < n and not _normalized(paragraphs[j].text):
            j += 1
        if j >= n:
            continue
        first = _normalized(paragraphs[j].text)
        if not first.startswith(('"', '“', '«')):
            continue
        # Protect until a closing quote appears or a conservative maximum is reached.
        for k in range(j, min(n, j + 14)):
            txt = _normalized(paragraphs[k].text)
            if not txt:
                break
            protected.add(id(paragraphs[k]._p))
            if k > j and re.search(r'[\"”»]\s*$', txt):
                break
            if k == j and len(txt) > 1 and re.search(r'[\"”»]\s*$', txt[1:]):
                break
    return protected


def _should_merge(
    current: Paragraph,
    nxt: Paragraph,
    frequencies: Counter,
    protected_nodes: Set[int],
) -> bool:
    a = _normalized(current.text)
    b = _normalized(nxt.text)

    if not a or not b:
        return False
    if id(current._p) in protected_nodes or id(nxt._p) in protected_nodes:
        return False
    if _is_page_number(a) or _is_page_number(b):
        return False
    if _looks_like_heading(a) or _looks_like_heading(b):
        return False

    # Repeated short body lines are usually copied headers/footers or recurring labels.
    if len(a) < 85 and frequencies.get(a, 0) >= 3:
        return False
    if len(b) < 85 and frequencies.get(b, 0) >= 3:
        return False

    if _ends_terminal(a):
        return False

    # A colon often introduces a quote/list/definition and should remain a paragraph break.
    if a.endswith(':'):
        return False

    # Conservative list/verse protection.
    if BULLET_RE.match(a) or BULLET_RE.match(b):
        return False
    if _short_line(a) and _short_line(b) and _starts_uppercase(b):
        return False

    # Strong continuation signals.
    if a.endswith((',', ';')):
        return True
    if _opening_balance(a) > 0:
        return True
    if _starts_lowercase(b):
        return len(a) >= 24

    # Proper nouns can begin with capitals in the middle of a sentence. Only join when the
    # first line is clearly a wrapped prose line rather than a short verse/title line.
    if _starts_uppercase(b) and len(a) >= 82 and len(b) >= 24:
        if not a.endswith(('"', '”', '»')):
            return True

    return False


def _merge_paragraphs(current: Paragraph, nxt: Paragraph) -> None:
    """Merge nxt into current while preserving run-level formatting as much as possible."""
    if current.text and not current.text.endswith((' ', '\t')):
        current.add_run(' ')

    nxt_el = nxt._p
    # Move runs/hyperlinks/bookmarks/etc., but not paragraph properties.
    for child in list(nxt_el):
        if child.tag.endswith('}pPr'):
            continue
        nxt_el.remove(child)
        current._p.append(child)

    parent = nxt_el.getparent()
    if parent is not None:
        parent.remove(nxt_el)


def repair_paragraph_sequence(paragraphs_getter) -> int:
    """Repair one paragraph sequence. paragraphs_getter must return a fresh list each call."""
    initial = list(paragraphs_getter())
    frequencies = _paragraph_frequencies(initial)
    protected_nodes = _protected_poetry_nodes(initial)
    repaired = 0
    i = 0

    while True:
        paragraphs = list(paragraphs_getter())
        if i >= len(paragraphs) - 1:
            break
        cur, nxt = paragraphs[i], paragraphs[i + 1]
        if _should_merge(cur, nxt, frequencies, protected_nodes):
            _merge_paragraphs(cur, nxt)
            repaired += 1
            # Keep i: the newly extended paragraph may need another continuation joined.
        else:
            i += 1
    return repaired


def _repair_document(doc: Document) -> int:
    repaired = repair_paragraph_sequence(lambda: doc.paragraphs)

    # Also handle text in table cells, without touching headers/footers.
    seen_cells = set()
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                key = id(cell._tc)
                if key in seen_cells:
                    continue
                seen_cells.add(key)
                repaired += repair_paragraph_sequence(lambda c=cell: c.paragraphs)
    return repaired


def output_path_for(source: Path) -> Path:
    return source.with_name(f"{source.stem}_Düzelti{source.suffix}")


def repair_file(source: os.PathLike | str) -> RepairResult:
    source = Path(source).expanduser().resolve()
    if source.suffix.lower() != '.docx':
        return RepairResult(source, source, 0, True, 'Yalnızca .docx dosyaları destekleniyor.')
    if source.stem.endswith('_Düzelti'):
        return RepairResult(source, source, 0, True, 'Bu dosya zaten _Düzelti kopyası gibi görünüyor.')
    if not source.exists():
        return RepairResult(source, source, 0, True, 'Dosya bulunamadı.')

    output = output_path_for(source)
    tmp = source.with_name(f".{source.stem}.{uuid.uuid4().hex}.tmp.docx")

    try:
        # Hard safety rule: never edit the selected original. Work on a physical copy.
        shutil.copy2(source, tmp)
        doc = Document(tmp)
        repaired = _repair_document(doc)
        doc.save(tmp)
        os.replace(tmp, output)
        return RepairResult(source, output, repaired, False, 'Tamamlandı.')
    except Exception as exc:
        try:
            if tmp.exists():
                tmp.unlink()
        except Exception:
            pass
        return RepairResult(source, output, 0, True, f'Hata: {exc}')


def collect_docx_from_folder(folder: os.PathLike | str, recursive: bool = False) -> List[Path]:
    folder = Path(folder).expanduser().resolve()
    pattern = '**/*.docx' if recursive else '*.docx'
    files = []
    for p in folder.glob(pattern):
        if p.is_file() and not p.name.startswith('~$') and not p.stem.endswith('_Düzelti'):
            files.append(p)
    return sorted(files, key=lambda p: str(p).lower())
