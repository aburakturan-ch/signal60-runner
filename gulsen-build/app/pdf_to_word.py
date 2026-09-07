from __future__ import annotations

import io
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import List

import pymupdf as fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageOps

from ocr_engine import ocr_image

SUPPORTED_IMAGE_EXTS = {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp', '.webp'}


@dataclass
class ConvertResult:
    source: Path
    output: Path
    pages: int
    images: int
    ocr_pages: int = 0
    skipped: bool = False
    message: str = ''
    cancelled: bool = False


def _cancelled(cancel_event) -> bool:
    try:
        return bool(cancel_event and cancel_event.is_set())
    except Exception:
        return False


def output_path_for_media(source: Path) -> Path:
    return source.with_name(f"{source.stem}_Aktarma.docx")


def collect_media_from_folder(folder: os.PathLike | str, recursive: bool = False) -> List[Path]:
    folder = Path(folder).expanduser().resolve()
    out = []
    it = folder.rglob('*') if recursive else folder.glob('*')
    for p in it:
        if not p.is_file() or p.name.startswith('~$'):
            continue
        suf = p.suffix.lower()
        if suf == '.pdf' or suf in SUPPORTED_IMAGE_EXTS:
            out.append(p)
    return sorted(out, key=lambda p: str(p).lower())


def _set_default_margins(doc: Document):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)


def _apply_style(run, span):
    size = span.get('size')
    if size:
        run.font.size = Pt(max(8, min(18, float(size))))
    flags = int(span.get('flags', 0) or 0)
    run.font.italic = bool(flags & 2)
    run.font.bold = bool(flags & 16)
    color = span.get('color')
    if isinstance(color, int):
        run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)


def _block_alignment(block, page_width):
    bbox = block.get('bbox', [0, 0, page_width, 0])
    x0, x1 = bbox[0], bbox[2]
    left = x0 / page_width
    right = (page_width - x1) / page_width
    if left > 0.22 and right > 0.22:
        return WD_ALIGN_PARAGRAPH.CENTER
    if left > 0.55 and right < 0.13:
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT


def _add_text_block(word: Document, block, page_width):
    lines = []
    for line in block.get('lines', []):
        spans = [s for s in line.get('spans', []) if s.get('text', '').strip()]
        if spans:
            lines.append(spans)
    if not lines:
        return False

    p = word.add_paragraph()
    p.alignment = _block_alignment(block, page_width)
    for li, line in enumerate(lines):
        if li:
            p.add_run(' ')
        prev = ''
        for span in line:
            text = re.sub(r'\s+', ' ', span.get('text', '')).strip()
            if not text:
                continue
            if prev and not prev.endswith((' ', '-', '—', '/', '(')) and not text.startswith(('.', ',', ';', ':', '!', '?', ')', '”', '"')):
                p.add_run(' ')
            r = p.add_run(text)
            _apply_style(r, span)
            prev = text
    return bool(p.text.strip())


def _save_block_image(block, tmpdir: Path) -> Path | None:
    data = block.get('image')
    if not data:
        return None
    ext = block.get('ext', 'png')
    path = tmpdir / f"image_{abs(hash(block.get('bbox', ())))}.{ext}"
    try:
        path.write_bytes(data)
        return path
    except Exception:
        return None


def _add_picture_safe(word: Document, path: Path, width=5.8):
    try:
        word.add_picture(str(path), width=Inches(width))
        return True
    except Exception:
        try:
            with Image.open(path) as im:
                im = ImageOps.exif_transpose(im).convert('RGB')
                tmp = path.with_suffix('.normalized.jpg')
                im.save(tmp, 'JPEG', quality=92)
            word.add_picture(str(tmp), width=Inches(width))
            try:
                tmp.unlink()
            except Exception:
                pass
            return True
        except Exception:
            return False


def _add_ocr_text(word: Document, text: str):
    text = text.strip()
    if not text:
        word.add_paragraph('[OCR ile okunabilir metin bulunamadı]')
        return
    for para in re.split(r'\n\s*\n', text):
        para = re.sub(r'[ \t]+', ' ', para).strip()
        if para:
            word.add_paragraph(para)


def _render_page_image(page) -> Image.Image:
    mat = fitz.Matrix(2.4, 2.4)
    pix = page.get_pixmap(matrix=mat, alpha=False)
    return Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB')


def _page_has_useful_text(page) -> bool:
    text = re.sub(r'\s+', ' ', page.get_text('text') or '').strip()
    return len(text) >= 35 and sum(ch.isalpha() for ch in text) >= 15


def _cancel_result(source: Path, out: Path, pages=0, imgs=0, ocr_pages=0):
    return ConvertResult(source, out, pages, imgs, ocr_pages, True, 'İptal edildi.', True)


def _convert_pdf(source: Path, include_images: bool, cancel_event=None) -> ConvertResult:
    out = output_path_for_media(source)
    word = Document()
    _set_default_margins(word)
    pages = imgs = ocr_pages = 0

    try:
        pdf = fitz.open(source)
        with tempfile.TemporaryDirectory(prefix='gulsen_pdf_') as td:
            tmpdir = Path(td)
            for pi, page in enumerate(pdf):
                if _cancelled(cancel_event):
                    pdf.close()
                    return _cancel_result(source, out, pages, imgs, ocr_pages)

                pages += 1
                if _page_has_useful_text(page):
                    pdata = page.get_text('dict', flags=11)
                    blocks = sorted(
                        pdata.get('blocks', []),
                        key=lambda b: (round(b.get('bbox', [0, 0, 0, 0])[1], 1), round(b.get('bbox', [0, 0, 0, 0])[0], 1)),
                    )
                    any_text = False
                    for block in blocks:
                        if _cancelled(cancel_event):
                            pdf.close()
                            return _cancel_result(source, out, pages, imgs, ocr_pages)
                        typ = block.get('type', 0)
                        if typ == 0:
                            any_text = _add_text_block(word, block, page.rect.width) or any_text
                        elif typ == 1 and include_images:
                            pth = _save_block_image(block, tmpdir)
                            if pth and _add_picture_safe(word, pth):
                                imgs += 1
                    if not any_text:
                        image = _render_page_image(page)
                        res = ocr_image(image, cancel_event=cancel_event)
                        if res.cancelled or _cancelled(cancel_event):
                            pdf.close()
                            return _cancel_result(source, out, pages, imgs, ocr_pages)
                        _add_ocr_text(word, res.text)
                        ocr_pages += 1
                else:
                    image = _render_page_image(page)
                    res = ocr_image(image, cancel_event=cancel_event)
                    if res.cancelled or _cancelled(cancel_event):
                        pdf.close()
                        return _cancel_result(source, out, pages, imgs, ocr_pages)
                    _add_ocr_text(word, res.text)
                    ocr_pages += 1
                    if include_images:
                        pth = tmpdir / f'page_{pi + 1}.jpg'
                        image.save(pth, 'JPEG', quality=88)
                        if _add_picture_safe(word, pth):
                            imgs += 1

                if pi < len(pdf) - 1:
                    word.add_page_break()

        pdf.close()
        if _cancelled(cancel_event):
            return _cancel_result(source, out, pages, imgs, ocr_pages)
        word.save(out)
        return ConvertResult(source, out, pages, imgs, ocr_pages, False, 'Tamamlandı.')
    except Exception as exc:
        if _cancelled(cancel_event):
            return _cancel_result(source, out, pages, imgs, ocr_pages)
        return ConvertResult(source, out, 0, 0, 0, True, f'Hata: {exc}')


def _convert_image(source: Path, include_images: bool, cancel_event=None) -> ConvertResult:
    out = output_path_for_media(source)
    word = Document()
    _set_default_margins(word)

    try:
        if _cancelled(cancel_event):
            return _cancel_result(source, out)
        with Image.open(source) as im:
            image = ImageOps.exif_transpose(im).convert('RGB')
        if _cancelled(cancel_event):
            return _cancel_result(source, out)

        res = ocr_image(image, deskew=True, cancel_event=cancel_event)
        if res.cancelled or _cancelled(cancel_event):
            return _cancel_result(source, out, 1, 0, 0)
        _add_ocr_text(word, res.text)

        imgs = 0
        if include_images and not _cancelled(cancel_event):
            with tempfile.TemporaryDirectory(prefix='gulsen_img_') as td:
                p = Path(td) / 'source.jpg'
                thumb = image.copy()
                thumb.thumbnail((2600, 2600), Image.Resampling.LANCZOS)
                thumb.save(p, 'JPEG', quality=90)
                if _add_picture_safe(word, p):
                    imgs = 1

        if _cancelled(cancel_event):
            return _cancel_result(source, out, 1, imgs, 1)
        word.save(out)
        return ConvertResult(source, out, 1, imgs, 1, False, 'Tamamlandı.')
    except Exception as exc:
        if _cancelled(cancel_event):
            return _cancel_result(source, out)
        return ConvertResult(source, out, 0, 0, 0, True, f'Hata: {exc}')


def convert_media_to_word(source: os.PathLike | str, include_images: bool = False, cancel_event=None) -> ConvertResult:
    source = Path(source).expanduser().resolve()
    if not source.exists():
        return ConvertResult(source, output_path_for_media(source), 0, 0, 0, True, 'Dosya bulunamadı.')
    if source.stem.endswith('_Aktarma'):
        return ConvertResult(source, output_path_for_media(source), 0, 0, 0, True, 'Bu dosya zaten _Aktarma çıktısı gibi görünüyor.')
    if _cancelled(cancel_event):
        return _cancel_result(source, output_path_for_media(source))
    suf = source.suffix.lower()
    if suf == '.pdf':
        return _convert_pdf(source, include_images, cancel_event=cancel_event)
    if suf in SUPPORTED_IMAGE_EXTS:
        return _convert_image(source, include_images, cancel_event=cancel_event)
    return ConvertResult(source, output_path_for_media(source), 0, 0, 0, True, 'Desteklenmeyen dosya türü.')


def collect_pdf_from_folder(folder, recursive=False):
    return [p for p in collect_media_from_folder(folder, recursive) if p.suffix.lower() == '.pdf']


def convert_pdf_to_word(source, include_images=False, cancel_event=None):
    return convert_media_to_word(source, include_images, cancel_event=cancel_event)
